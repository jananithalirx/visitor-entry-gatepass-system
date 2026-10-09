from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
import os
from datetime import datetime

app = Flask(__name__)

EXCEL_FILE = "visitors_data.xlsx"


# checking if file exists, if not create it with headers
def create_excel_if_not_exists():
    if not os.path.exists(EXCEL_FILE):
        print("creating new excel file...")
        wb = Workbook()
        ws = wb.active
        ws.title = "Visitors"
        ws.append(["ID", "Name", "Phone", "Purpose", "Meeting", "Vehicle No", "In Time", "Out Time", "Date"])
        wb.save(EXCEL_FILE)
        print("excel file created with headers")
    else:
        print("excel file already exists, skipping creation")


create_excel_if_not_exists()


def get_next_id():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    # counting rows minus the header row
    last_id = ws.max_row - 1
    return last_id + 1


@app.route('/')
def home():
    return render_template('index.html')


@app.route('/new_entry', methods=['GET', 'POST'])
def new_entry():
    if request.method == 'POST':
        v_name = request.form.get('v_name')
        ph_no = request.form.get('ph_no')
        purpose = request.form.get('purpose')
        meet_person = request.form.get('meet_person')
        vehicle_no = request.form.get('vehicle_no')

        # TODO: add strict validation later, for now just a basic check
        if v_name == "" or ph_no == "":
            return render_template('new_entry.html', error="Name and Phone are required!")

        wb = load_workbook(EXCEL_FILE)
        ws = wb.active

        v_id = get_next_id()
        in_time = datetime.now().strftime("%I:%M %p")
        today_date = datetime.now().strftime("%d-%m-%Y")

        ws.append([v_id, v_name, ph_no, purpose, meet_person, vehicle_no, in_time, "", today_date])
        wb.save(EXCEL_FILE)
        print("Data saved successfully, visitor id =", v_id)

        return redirect(url_for('gatepass', v_id=v_id))

    return render_template('new_entry.html', error=None)


@app.route('/gatepass/<int:v_id>')
def gatepass(v_id):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    visitor = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == v_id:
            visitor = row
            break
    if visitor is None:
        return "Visitor not found", 404
    data = {
        "id": visitor[0],
        "name": visitor[1],
        "phone": visitor[2],
        "purpose": visitor[3],
        "meet": visitor[4],
        "vehicle": visitor[5],
        "in_time": visitor[6],
        "date": visitor[8]
    }
    return render_template('gatepass.html', v=data)
@app.route('/records')
def records():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    all_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        all_rows.append(row)

    # reversing so latest entry shows on top
    all_rows.reverse()

    return render_template('records.html', rows=all_rows)


@app.route('/checkout', methods=['GET', 'POST'])
def checkout():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    if request.method == 'POST':
        v_id = int(request.form.get('v_id'))
        out_time = datetime.now().strftime("%I:%M %p")

        # loop through rows, find the matching id, update the out time col
        for row in ws.iter_rows(min_row=2):
            if row[0].value == v_id:
                row[7].value = out_time  # out time column
                print("checkout done for id", v_id)
                break

        wb.save(EXCEL_FILE)
        return redirect(url_for('checkout'))

    # for GET request, show list of people who haven't checked out yet
    pending = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[7] == "" or row[7] is None:
            pending.append(row)

    return render_template('checkout.html', pending=pending)


if __name__ == '__main__':
    app.run(debug=True)
